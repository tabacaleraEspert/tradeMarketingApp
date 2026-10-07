"""Migración de datos 2026-10-07: PdvSupplier legacy (texto Name/Phone) → Supplier + SupplierSeller.

Requiere el DDL de 0026 aplicado (local: `alembic upgrade head`; prod:
scripts/hotfix_supplier_sellers_prod_20261007.py). Conexión = app.database
(env-driven: sin DATABASE_USER/PASSWORD → SQLite local; con ellas → Azure SQL).

Dos pasos, con revisión humana en el medio:

  1. Propuesta (solo lectura):
       python scripts/migrate_suppliers_sellers.py --propose propuesta.xlsx
     Por zona, agrupa las filas PdvSupplier sin SupplierId por nombre normalizado
     (minúsculas, sin tildes, espacios colapsados; para agrupar además k→c) y
     fusiona por similitud (difflib ratio >= 0.85) dentro de la zona. Una fila
     por (zona, nombre original, teléfono). Columnas editables: ProposedSupplier
     (vacío = no migrar esa fila) y SellerName. Teléfonos con < 8 dígitos se
     marcan y se descartan (vendedor sin teléfono).
     Sin openpyxl escribe CSV.

  2. Aplicar (escribe; backup JSON previo de las filas afectadas):
       python scripts/migrate_suppliers_sellers.py --apply propuesta.xlsx [--dry-run]
     Crea un Supplier por (zona, ProposedSupplier) — reusa uno activo con mismo
     nombre normalizado —, un SupplierSeller por teléfono válido distinto (nombre =
     SellerName o "Vendedor <tel>"; sin teléfono válido ni SellerName → sin
     vendedor) y vincula PdvSupplier.SupplierId/SupplierSellerId (+ columnas
     legacy). Tipo = el más común de sus filas; productos = unión.
     Idempotente: las filas que ya tienen SupplierId se saltean.
     Si un PDV queda con 2 filas activas del mismo proveedor, la segunda se desactiva.
"""
from __future__ import annotations

import argparse
import csv
import difflib
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.database import SessionLocal  # noqa: E402
from app.models import PDV, PdvSupplier, Zone  # noqa: E402
from app.models.supplier import Supplier, SupplierSeller  # noqa: E402
from app.services.supplier_names import normalize_name  # noqa: E402
from app.services.suppliers import (  # noqa: E402
    audit,
    find_active_supplier_by_name,
    json_to_products,
    link_snapshot,
    products_to_json,
    supplier_snapshot,
    sync_link_legacy,
)

COLUMNS = ["ZoneId", "ZoneName", "RowName", "Phone", "PdvCount", "ProposedSupplier", "SellerName", "Notes"]
FUZZY_RATIO = 0.85
MIN_PHONE_DIGITS = 8


# ─── Helpers ─────────────────────────────────────────────────────────

def group_key(name: str) -> str:
    """Clave de agrupación: normalize_name + k→c ("Distri Kosiuko" ~ "Distri Cosiuko")."""
    return normalize_name(name).replace("k", "c")


def phone_digits(phone: str | None) -> str:
    return re.sub(r"\D", "", phone or "")


def valid_phone(phone: str | None) -> str | None:
    """Teléfono limpio si tiene >= 8 dígitos; si no, None (se descarta)."""
    p = (phone or "").strip()
    return p if len(phone_digits(p)) >= MIN_PHONE_DIGITS else None


def _resolved_zone(row: PdvSupplier, pdv_zone: dict[int, int | None]) -> int | None:
    """Zona del PDV y, si no tiene, la de la fila (misma regla que POST /link)."""
    z = pdv_zone.get(row.PdvId) if row.PdvId is not None else None
    return z if z is not None else row.ZoneId


def _pending_rows(db):
    """Filas activas sin vincular + zona resuelta (fila → PDV)."""
    rows = (
        db.query(PdvSupplier)
        .filter(PdvSupplier.SupplierId.is_(None), PdvSupplier.IsActive == True)  # noqa: E712
        .order_by(PdvSupplier.PdvSupplierId)
        .all()
    )
    pdv_ids = {r.PdvId for r in rows if r.PdvId is not None}
    pdv_zone = (
        {p.PdvId: p.ZoneId for p in db.query(PDV.PdvId, PDV.ZoneId).filter(PDV.PdvId.in_(pdv_ids)).all()}
        if pdv_ids else {}
    )
    return [(r, _resolved_zone(r, pdv_zone)) for r in rows]


def _row_key(zone_id, name, phone) -> tuple:
    return (zone_id, (name or "").strip(), (phone or "").strip())


# ─── Propuesta ───────────────────────────────────────────────────────

# Palabras genéricas: "LAG distribuciones" ~ "Tato distribuciones" da ratio alto
# solo por el sufijo; la parte distintiva ("lag" vs "tato") tiene que coincidir.
GENERIC_WORDS = {
    "distribuciones", "distribucion", "distribuidora", "distribuidor", "distri",
    "cigarrera", "cigarreria", "ciosco", "ciosc", "mayorista", "intermediario", "intermediarios",
    "sa", "srl", "de", "del", "la", "el", "los", "las", "y", "av", "avenida",
}


def _distinct_part(key: str) -> str:
    return " ".join(w for w in key.replace(".", " ").split() if w not in GENERIC_WORDS)


def _distinct_parts_match(a: str, b: str) -> bool:
    da, db_ = _distinct_part(a), _distinct_part(b)
    if not da or not db_:
        return True
    if da in db_ or db_ in da:
        return True
    return difflib.SequenceMatcher(None, da, db_).ratio() >= FUZZY_RATIO


def cluster_keys(key_counts: Counter) -> dict[str, str]:
    """Agrupa claves similares (ratio >= FUZZY_RATIO). Greedy: las más frecuentes
    primero son representantes. Devuelve clave → clave representante."""
    reps: list[str] = []
    mapping: dict[str, str] = {}
    for key, _ in sorted(key_counts.items(), key=lambda kv: (-kv[1], kv[0])):
        best, best_ratio = None, 0.0
        for rep in reps:
            ratio = difflib.SequenceMatcher(None, key, rep).ratio()
            if ratio >= FUZZY_RATIO and ratio > best_ratio and _distinct_parts_match(key, rep):
                best, best_ratio = rep, ratio
        if best is None:
            reps.append(key)
            mapping[key] = key
        else:
            mapping[key] = best
    return mapping


def build_proposal(db) -> list[dict]:
    zones = {z.ZoneId: z.Name for z in db.query(Zone).all()}
    by_zone: dict[int | None, list[PdvSupplier]] = defaultdict(list)
    for row, zone_id in _pending_rows(db):
        by_zone[zone_id].append(row)

    out: list[dict] = []
    for zone_id, rows in sorted(by_zone.items(), key=lambda kv: zones.get(kv[0], "~")):
        key_counts = Counter(group_key(r.Name) for r in rows)
        rep_of = cluster_keys(key_counts)

        # Nombre propuesto por cluster: el nombre original más frecuente
        names_by_rep: dict[str, Counter] = defaultdict(Counter)
        for r in rows:
            names_by_rep[rep_of[group_key(r.Name)]][r.Name.strip()] += 1
        proposed = {rep: c.most_common(1)[0][0] for rep, c in names_by_rep.items()}

        # Una línea por (nombre original, teléfono)
        lines: dict[tuple, dict] = {}
        for r in rows:
            k = _row_key(zone_id, r.Name, r.Phone)
            line = lines.get(k)
            if line is None:
                rep = rep_of[group_key(r.Name)]
                notes = []
                if (r.Phone or "").strip() and not valid_phone(r.Phone):
                    notes.append(f"teléfono inválido (<{MIN_PHONE_DIGITS} dígitos): se descarta")
                if rep != group_key(r.Name):
                    notes.append(f"fusionado por similitud con «{proposed[rep]}»")
                line = lines[k] = {
                    "ZoneId": zone_id if zone_id is not None else "",
                    "ZoneName": zones.get(zone_id, "") if zone_id is not None else "",
                    "RowName": k[1],
                    "Phone": k[2],
                    "_pdvs": set(),
                    "ProposedSupplier": proposed[rep],
                    "SellerName": "",
                    "Notes": "; ".join(notes),
                }
            if r.PdvId is not None:
                line["_pdvs"].add(r.PdvId)
        for line in sorted(lines.values(), key=lambda x: (x["ProposedSupplier"].lower(), x["RowName"].lower(), x["Phone"])):
            line["PdvCount"] = len(line.pop("_pdvs"))
            out.append(line)
    return out


def write_mapping(lines: list[dict], path: Path) -> Path:
    if path.suffix.lower() == ".xlsx":
        try:
            from openpyxl import Workbook
        except ImportError:
            path = path.with_suffix(".csv")
            print(f"openpyxl no disponible — escribo CSV: {path}")
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = "Propuesta"
            ws.append(COLUMNS)
            for line in lines:
                ws.append([line[c] for c in COLUMNS])
            ws.freeze_panes = "A2"
            wb.save(path)
            return path
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS)
        w.writeheader()
        for line in lines:
            w.writerow({c: line[c] for c in COLUMNS})
    return path


def read_mapping(path: Path) -> list[dict]:
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True, data_only=True).active
        it = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(it)]
        rows = [dict(zip(header, vals)) for vals in it]
    else:
        with path.open(newline="", encoding="utf-8-sig") as f:
            rows = list(csv.DictReader(f))

    def s(v) -> str:
        if v is None:
            return ""
        if isinstance(v, float) and v.is_integer():
            v = int(v)
        return str(v).strip()

    return [{k: s(v) for k, v in r.items()} for r in rows if any(v not in (None, "") for v in r.values())]


# ─── Aplicar ─────────────────────────────────────────────────────────

def apply_mapping(db, mapping: list[dict], dry_run: bool = False, backup_dir: Path | None = None) -> dict:
    """Aplica el mapeo. Sin commit si dry_run (rollback al final). Devuelve resumen."""
    zones_by_name = {z.Name: z.ZoneId for z in db.query(Zone).all()}

    # (zona, nombre original, teléfono) → (ProposedSupplier, SellerName)
    plan: dict[tuple, tuple[str, str]] = {}
    for m in mapping:
        proposed = m.get("ProposedSupplier", "")
        if not proposed:
            continue
        zid_raw = m.get("ZoneId", "")
        if zid_raw:
            zone_id = int(zid_raw)
        elif m.get("ZoneName"):
            zone_id = zones_by_name.get(m["ZoneName"])
            if zone_id is None:
                print(f"  ! zona desconocida «{m['ZoneName']}» — salto «{m.get('RowName')}»")
                continue
        else:
            zone_id = None
        plan[_row_key(zone_id, m.get("RowName"), m.get("Phone"))] = (proposed, m.get("SellerName", ""))

    affected = [(r, z) for r, z in _pending_rows(db) if _row_key(z, r.Name, r.Phone) in plan]
    summary = {
        "filasMapeo": len(mapping), "filasAfectadas": len(affected),
        "proveedoresCreados": 0, "proveedoresReusados": 0, "vendedoresCreados": 0,
        "vinculadas": 0, "duplicadasDesactivadas": 0, "backup": None,
    }
    if not affected:
        return summary

    # Backup: estado previo de toda fila que se toca (afectadas + duplicadas ya
    # vinculadas que reciben vendedor). Se escribe antes del commit.
    def _row_backup(r: PdvSupplier) -> dict:
        return {c.name: getattr(r, c.name) for c in PdvSupplier.__table__.columns}

    backup_rows: dict[int, dict] = {r.PdvSupplierId: _row_backup(r) for r, _ in affected}
    backup_path = None
    if not dry_run:
        backup_dir = backup_dir or Path(__file__).resolve().parent
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup_path = backup_dir / f"backup_migrate_suppliers_{ts}.json"
        summary["backup"] = str(backup_path)

    # Agrupar por (zona, proveedor normalizado)
    groups: dict[tuple, list] = defaultdict(list)
    display: dict[tuple, str] = {}
    for r, z in affected:
        proposed, seller_name = plan[_row_key(z, r.Name, r.Phone)]
        gk = (z, normalize_name(proposed))
        groups[gk].append((r, seller_name))
        display.setdefault(gk, proposed.strip())

    try:
        for (zone_id, _), items in groups.items():
            name = display[(zone_id, _)]
            types = Counter(r.SupplierTypeId for r, _s in items if r.SupplierTypeId)
            products: set[str] = set()
            for r, _s in items:
                products.update(json_to_products(r.Products) or [])

            supplier = find_active_supplier_by_name(db, zone_id, name)
            if supplier is None:
                supplier = Supplier(
                    ZoneId=zone_id, Name=name,
                    SupplierTypeId=types.most_common(1)[0][0] if types else None,
                    Products=products_to_json(sorted(products)) if products else None,
                    IsActive=True,
                )
                db.add(supplier)
                db.flush()
                audit(db, None, "Supplier", supplier.SupplierId, "SUPPLIER_CREATE",
                      {"despues": supplier_snapshot(supplier), "origen": "migrate_suppliers_sellers"})
                summary["proveedoresCreados"] += 1
            else:
                if supplier.SupplierTypeId is None and types:
                    supplier.SupplierTypeId = types.most_common(1)[0][0]
                merged = set(json_to_products(supplier.Products) or []) | products
                if merged:
                    supplier.Products = products_to_json(sorted(merged))
                summary["proveedoresReusados"] += 1

            # Vendedores existentes del proveedor (por teléfono y por nombre)
            existing = db.query(SupplierSeller).filter(
                SupplierSeller.SupplierId == supplier.SupplierId, SupplierSeller.IsActive == True,  # noqa: E712
            ).all()
            by_phone = {phone_digits(s.Phone): s for s in existing if s.Phone}
            by_name = {normalize_name(s.Name): s for s in existing}

            linked_pdvs: dict[int, PdvSupplier] = {
                r.PdvId: r for r in db.query(PdvSupplier).filter(
                    PdvSupplier.SupplierId == supplier.SupplierId, PdvSupplier.IsActive == True,  # noqa: E712
                    PdvSupplier.PdvId.isnot(None),
                ).all()
            }

            for r, seller_name in items:
                phone = valid_phone(r.Phone)
                seller = None
                if phone:
                    seller = by_phone.get(phone_digits(phone))
                    if seller is None:
                        sname = seller_name or f"Vendedor {phone}"
                        seller = by_name.get(normalize_name(sname))
                        if seller is not None and not seller.Phone:
                            seller.Phone = phone
                            by_phone[phone_digits(phone)] = seller
                elif seller_name:
                    seller = by_name.get(normalize_name(seller_name))
                if seller is None and (phone or seller_name):
                    sname = seller_name or f"Vendedor {phone}"
                    seller = SupplierSeller(SupplierId=supplier.SupplierId, Name=sname, Phone=phone, IsActive=True)
                    db.add(seller)
                    db.flush()
                    summary["vendedoresCreados"] += 1
                    by_name[normalize_name(sname)] = seller
                    if phone:
                        by_phone[phone_digits(phone)] = seller

                before = link_snapshot(r)
                dup = linked_pdvs.get(r.PdvId) if r.PdvId is not None else None
                if dup is not None and dup.PdvSupplierId != r.PdvSupplierId:
                    if dup.SupplierSellerId is None and seller is not None:
                        backup_rows.setdefault(dup.PdvSupplierId, _row_backup(dup))
                        dup_before = link_snapshot(dup)
                        sync_link_legacy(dup, supplier, seller)
                        db.flush()
                        audit(db, None, "PdvSupplier", dup.PdvSupplierId, "PDV_SUPPLIER_LINK",
                              {"antes": dup_before, "despues": link_snapshot(dup),
                               "origen": "migrate_suppliers_sellers"})
                    r.SupplierId = supplier.SupplierId
                    r.IsActive = False
                    summary["duplicadasDesactivadas"] += 1
                else:
                    sync_link_legacy(r, supplier, seller)
                    if r.PdvId is not None:
                        linked_pdvs[r.PdvId] = r
                    summary["vinculadas"] += 1
                db.flush()
                audit(db, None, "PdvSupplier", r.PdvSupplierId, "PDV_SUPPLIER_LINK",
                      {"antes": before, "despues": link_snapshot(r), "origen": "migrate_suppliers_sellers"})

        if dry_run:
            db.rollback()
        else:
            backup_path.write_text(json.dumps(
                list(backup_rows.values()), default=str, ensure_ascii=False, indent=1,
            ), encoding="utf-8")
            print(f"Backup: {backup_path}")
            db.commit()
    except Exception:
        db.rollback()
        raise
    return summary


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--propose", metavar="OUT.xlsx|csv")
    g.add_argument("--apply", metavar="MAPPING.xlsx|csv")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    from app.database import engine
    print(f"DB: {engine.url.render_as_string(hide_password=True)}")
    db = SessionLocal()
    try:
        if args.propose:
            lines = build_proposal(db)
            path = write_mapping(lines, Path(args.propose))
            n_sup = len({(l["ZoneId"], normalize_name(l["ProposedSupplier"])) for l in lines})
            print(f"Propuesta: {len(lines)} líneas → {n_sup} proveedores propuestos. Archivo: {path}")
        else:
            mapping = read_mapping(Path(args.apply))
            summary = apply_mapping(db, mapping, dry_run=args.dry_run)
            print(("[dry-run] " if args.dry_run else "") + json.dumps(summary, ensure_ascii=False, indent=1))
    finally:
        db.close()


if __name__ == "__main__":
    main()
